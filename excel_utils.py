# excel_utils.py

from openpyxl import Workbook
from openpyxl.chart import BarChart, PieChart, Reference
import pandas as pd

def generar_graficos_excel(kpi_data, clicks_data):
    wb = Workbook()
    ws = wb.active
    ws.title = "KPI y Clics"

    # Insertar KPI data
    ws.append(["Métrica", "Valor"])
    ws.append(["Total Records", kpi_data["TotalRecords"]])
    ws.append(["Complete Records", kpi_data["CompleteRecords"]])
    ws.append(["Correction Records", kpi_data["CorrectionRecords"]])

    # Gráfico de barras para KPI
    bar_chart = BarChart()
    data = Reference(ws, min_col=2, min_row=2, max_row=4)
    cats = Reference(ws, min_col=1, min_row=2, max_row=4)
    bar_chart.add_data(data, titles_from_data=False)
    bar_chart.set_categories(cats)
    bar_chart.title = "KPI Summary"
    ws.add_chart(bar_chart, "E5")

    # Insertar datos de clicks
    ws.append(["Fecha", "Button Name", "Click Count"])
    for index, row in clicks_data.iterrows():
        ws.append([row['ClickDate'], row['ButtonName'], row['ClickCount']])

    # Pie Chart for a summary of clicks per button (aggregate by button name)
    pie_chart = PieChart()
    clicks_summary = clicks_data.groupby('ButtonName').sum().reset_index()
    start_row = ws.max_row + 2
    ws.append(["Button Name", "Total Clicks"])
    for i, row in clicks_summary.iterrows():
        ws.append([row['ButtonName'], row['ClickCount']])

    pie_data = Reference(ws, min_col=2, min_row=start_row + 1, max_row=start_row + clicks_summary.shape[0])
    pie_labels = Reference(ws, min_col=1, min_row=start_row + 1, max_row=start_row + clicks_summary.shape[0])
    pie_chart.add_data(pie_data, titles_from_data=True)
    pie_chart.set_categories(pie_labels)
    pie_chart.title = "Total Clicks by Button"
    ws.add_chart(pie_chart, "E20")

    wb.save("kpi_clicks_report.xlsx")
